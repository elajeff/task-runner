from pathlib import Path
import json

import pytest

from runtime.documents import extract_document


def test_oversized_document_is_rejected(tmp_path: Path, monkeypatch):
    import runtime.documents as documents
    source = tmp_path / "big.txt"
    source.write_text("x" * 30)
    monkeypatch.setattr(documents, "MAX_DOCUMENT_BYTES", 10)
    with pytest.raises(ValueError, match="larger than"):
        extract_document(source)


def test_missing_ocr_dependency_has_install_instructions(tmp_path: Path, monkeypatch):
    import builtins
    import runtime.documents as documents
    source = tmp_path / "scan.png"
    source.write_bytes(b"not a real image")
    real_import = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name == "indra_ocr":
            raise ImportError("mock missing OCR")
        return real_import(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    with pytest.raises(RuntimeError, match="pip install"):
        documents.extract_document(source, use_ocr=True)


def test_long_text_is_split_into_bounded_ai_sections(tmp_path: Path):
    from runtime.documents import _split_sections, MAX_SECTION_CHARS
    source = tmp_path / "long.txt"
    source.write_text(("Paragraph content. " * 1200) + "\n\n" + ("Second paragraph. " * 1200), encoding="utf-8")
    document = extract_document(source)
    sections = _split_sections(document.pages, page_labels=False)
    assert len(sections) > 1
    assert all(len(section) <= MAX_SECTION_CHARS for section in sections)
    assert all("[Page 1]" not in section for section in sections)


def test_scanned_pdf_does_not_invoke_indra_when_ocr_is_disabled(tmp_path: Path, monkeypatch):
    import runtime.documents as documents

    source = tmp_path / "scan.pdf"
    source.write_bytes(b"mock pdf bytes")

    monkeypatch.setattr(documents, "_pdfplumber_extract_pdf", lambda _path: ([(1, "")], []))
    monkeypatch.setattr(documents, "_indra_extract", lambda _path: (_ for _ in ()).throw(AssertionError("OCR should not run")))
    with pytest.raises(ValueError, match="Use OCR for scanned PDFs and images"):
        documents.extract_document(source, use_ocr=False)


def test_document_job_saves_progress_and_skips_completed_sections_on_resume(tmp_path: Path):
    from runtime.documents import create_document_job, load_document_job, process_document_job

    source = tmp_path / "long-notes.txt"
    source.write_text("Useful lecture sentence. " * 2200, encoding="utf-8")
    job_dir = create_document_job(source, ["summary"], False, tmp_path / "document_outputs")
    prompts: list[str] = []

    def provider_fails_once(prompt: str) -> str:
        prompts.append(prompt)
        if len(prompts) == 2:
            raise RuntimeError("temporary provider failure")
        return json.dumps({"summary": f"Summary from call {len(prompts)}"})

    with pytest.raises(RuntimeError, match="temporary provider failure"):
        process_document_job(job_dir, provider_fails_once)

    saved_after_failure = load_document_job(job_dir)
    assert saved_after_failure["status"] == "failed"
    assert "1" in saved_after_failure["completed_sections"]
    assert saved_after_failure["last_completed_step"].startswith("Generated selected AI outputs for section 1")

    resumed_prompts: list[str] = []

    def provider_after_resume(prompt: str) -> str:
        resumed_prompts.append(prompt)
        return json.dumps({"summary": f"Resumed summary {len(resumed_prompts)}"})

    result_dir = process_document_job(job_dir, provider_after_resume)
    assert result_dir == job_dir
    assert load_document_job(job_dir)["status"] == "completed"
    assert len([p for p in prompts if "Source section 1/" in p]) == 1
    assert all("Source section 1/" not in prompt for prompt in resumed_prompts)
    assert (job_dir / "02_AI_Results" / "summary.md").is_file()
    assert "Resumed summary" in (job_dir / "02_AI_Results" / "summary.md").read_text(encoding="utf-8")


def test_ocr_warning_is_separate_from_exact_verbatim_extraction(tmp_path: Path):
    from runtime.documents import (
        ExtractedDocument, create_document_job, load_document_job, process_document_job,
    )

    source = tmp_path / "scan.png"
    source.write_bytes(b"mock image")
    job_dir = create_document_job(source, ["verbatim"], True, tmp_path / "document_outputs")

    def fake_extract(path: Path, *, use_ocr: bool = False) -> ExtractedDocument:
        assert use_ocr is True
        return ExtractedDocument(path, [(1, "Exact line.\nSecond line.")], ocr_used=True)

    process_document_job(job_dir, None, extract_fn=fake_extract)
    state = load_document_job(job_dir)
    assert state["status"] == "completed"
    assert (job_dir / "01_Extracted_Text" / "extracted_text.md").read_text(encoding="utf-8") == "Exact line.\nSecond line."
    notes = (job_dir / "03_Processing_Notes" / "processing_notes.md").read_text(encoding="utf-8")
    assert "OCR was used" in notes
    assert "OCR was used" not in (job_dir / "01_Extracted_Text" / "extracted_text.md").read_text(encoding="utf-8")
    assert not (job_dir / "START_HERE.txt").exists()
    assert (job_dir / "01_Extracted_Text" / "extracted_text.md").is_file()
    assert (job_dir / "03_Processing_Notes" / "processing_notes.md").is_file()


def test_resumable_job_list_excludes_completed_jobs(tmp_path: Path):
    from runtime.documents import create_document_job, list_resumable_document_jobs, process_document_job

    source = tmp_path / "notes.txt"
    source.write_text("Some notes.", encoding="utf-8")
    root = tmp_path / "outputs"
    waiting = create_document_job(source, ["verbatim"], False, root)
    completed = create_document_job(source, ["verbatim"], False, root)
    process_document_job(completed, None)
    jobs = list_resumable_document_jobs(root)
    ids = {Path(job["job_dir"]) for job in jobs}
    assert waiting in ids
    assert completed not in ids


def test_mixed_pdf_without_ocr_keeps_text_and_warns_about_unreadable_pages(tmp_path: Path, monkeypatch):
    import runtime.documents as documents

    source = tmp_path / "mixed.pdf"
    source.write_bytes(b"mock mixed pdf")

    monkeypatch.setattr(
        documents,
        "_pdfplumber_extract_pdf",
        lambda _path: ([(1, "Selectable lecture text"), (2, "")], []),
    )
    monkeypatch.setattr(documents, "_indra_extract", lambda _path: (_ for _ in ()).throw(AssertionError("OCR should not run")))
    document = documents.extract_document(source, use_ocr=False)
    assert "Selectable lecture text" in document.text
    assert document.ocr_used is False
    assert document.warnings and "pages 2" in document.warnings[0]



def test_table_aware_pdf_layout_is_used_when_ocr_is_also_enabled(tmp_path: Path, monkeypatch):
    import runtime.documents as documents

    source = tmp_path / "report.pdf"
    source.write_bytes(b"mock pdf")
    monkeypatch.setattr(documents, "_indra_extract", lambda _path: ([(1, "Indra text"), (2, "Scanned page text")], True))
    monkeypatch.setattr(
        documents,
        "_pdfplumber_extract_pdf",
        lambda _path: ([(1, "Report\n\n|Name|Count|\n|---|---|\n|Alpha|12|"), (2, "")], ["Table detected."]),
    )

    result = documents.extract_document(source, use_ocr=True)

    assert result.pages[0][1].startswith("Report")
    assert "|Alpha|12|" in result.pages[0][1]
    assert result.pages[1][1] == "Scanned page text"
    assert result.ocr_used is True
    assert result.warnings == ["Table detected."]

def test_indra_ocr_uses_documented_python_api_only_when_requested(tmp_path: Path, monkeypatch):
    import sys
    from types import SimpleNamespace
    import runtime.documents as documents

    source = tmp_path / "scan.pdf"
    source.write_bytes(b"mock pdf bytes")

    class FakePdfPage:
        def extract_text(self, **kwargs):
            return ""

    class FakeReader:
        pages = [FakePdfPage()]

    class OcrPage:
        text = "recognized text"
        markdown = "# recognized text\n\n中文文本"
        source = "ocr"

    calls = {}
    class Converter:
        def convert(self, path, max_pages):
            calls["convert"] = (path, max_pages)
            return [(1, OcrPage())]

    def load(**kwargs):
        calls["load"] = kwargs
        return Converter()

    monkeypatch.setitem(sys.modules, "pypdf", SimpleNamespace(PdfReader=lambda _path: FakeReader()))
    monkeypatch.setitem(sys.modules, "indra_ocr", SimpleNamespace(load=load))
    document = documents.extract_document(source, use_ocr=True)
    assert document.pages == [(1, "# recognized text\n\n中文文本")]
    assert document.ocr_used is True
    assert calls["load"] == {"threads": 2, "mode": "auto"}
    assert calls["convert"] == (str(source.resolve()), documents.MAX_DOCUMENT_PAGES)


def test_same_document_job_cannot_be_processed_concurrently(tmp_path: Path):
    import threading
    from runtime.documents import create_document_job, process_document_job

    source = tmp_path / "notes.txt"
    source.write_text("A useful concept.", encoding="utf-8")
    job_dir = create_document_job(source, ["summary"], False, tmp_path / "document_outputs")
    provider_started = threading.Event()
    release_provider = threading.Event()
    worker_errors: list[Exception] = []

    def slow_provider(_prompt: str) -> str:
        provider_started.set()
        if not release_provider.wait(timeout=3):
            raise TimeoutError("Test provider was not released.")
        return json.dumps({"summary": "Completed summary."})

    def run_job() -> None:
        try:
            process_document_job(job_dir, slow_provider)
        except Exception as exc:
            worker_errors.append(exc)

    worker = threading.Thread(target=run_job)
    worker.start()
    assert provider_started.wait(timeout=2)
    try:
        with pytest.raises(RuntimeError, match="already running"):
            process_document_job(job_dir, lambda _prompt: json.dumps({"summary": "Second run."}))
    finally:
        release_provider.set()
        worker.join(timeout=3)
    assert not worker.is_alive()
    assert worker_errors == []


def test_malformed_ai_response_leaves_a_recoverable_document_job(tmp_path: Path):
    from runtime.documents import create_document_job, list_resumable_document_jobs, load_document_job, process_document_job

    source = tmp_path / "notes.txt"
    source.write_text("Definitions need accurate sources.", encoding="utf-8")
    job_dir = create_document_job(source, ["definitions"], False, tmp_path / "document_outputs")
    with pytest.raises(RuntimeError, match="malformed output"):
        process_document_job(job_dir, lambda _prompt: "this is not JSON")
    saved = load_document_job(job_dir)
    assert saved["status"] == "failed"
    assert saved["pages"] == [[1, "Definitions need accurate sources."]]
    assert any(Path(job["job_dir"]) == job_dir for job in list_resumable_document_jobs(tmp_path / "document_outputs"))


def test_pdf_extraction_detects_and_preserves_tables():
    source = Path(__file__).parent / "fixtures" / "structured_table.pdf"
    document = extract_document(source)
    text = document.text
    assert "Quarterly revenue report" in text
    assert "| Product | Revenue | Year |" in text
    assert "| Alpha | 1,200 | 2025 |" in text
    assert "End of report." in text
    assert any("tables were detected" in warning for warning in document.warnings)


def test_borderless_pdf_table_detection_keeps_rows_and_values():
    import pdfplumber

    source = Path(__file__).parent / "fixtures" / "borderless_table.pdf"

    # Validate the fixture separately from the table formatter. If this fails,
    # the test PDF no longer contains the selectable caption this regression
    # test is intended to preserve; changing production extraction code would
    # not be a valid fix for missing source text.
    with pdfplumber.open(source) as pdf:
        page = pdf.pages[0]
        raw_text = page.extract_text() or ""
        positioned_lines = [line.get("text", "") for line in page.extract_text_lines()]

    assert "Figures are illustrative." in raw_text, (
        "The borderless-table fixture must contain 'Figures are illustrative.' "
        "as selectable text. Check the fixture, not the table formatter. "
        f"pdfplumber={pdfplumber.__version__}; extracted text={raw_text!r}"
    )

    document = extract_document(source)
    text = document.text
    assert "Products by year" in text
    assert "| Product | Revenue | Year |" in text
    assert "| Alpha | 1,200 | 2025 |" in text
    assert "| Beta | 2,350 | 2026 |" in text
    assert "Figures are illustrative." in text, (
        "The caption exists in the original PDF's extracted text but was lost "
        "during table-aware extraction. "
        f"pdfplumber={pdfplumber.__version__}; positioned lines={positioned_lines!r}; "
        f"raw text={raw_text!r}; formatted text={text!r}"
    )
    assert text.index("Products by year") < text.index("| Product |") < text.index("| Beta |") < text.index("Figures are illustrative.")
    assert not any(line.strip() == "Beta 2,350 2026" for line in text.splitlines())
    assert not any(line.strip() == "Product Revenue Year" for line in text.splitlines())
    assert any("tables were detected" in warning for warning in document.warnings)


def test_local_ocr_output_uses_markdown_extension_for_table_layout(tmp_path: Path):
    from runtime.documents import ExtractedDocument, create_document_job, process_document_job

    source = tmp_path / "scan.png"
    source.write_bytes(b"mock image")
    job_dir = create_document_job(source, ["verbatim"], True, tmp_path / "outputs")

    def fake_extract(path: Path, *, use_ocr: bool = False) -> ExtractedDocument:
        assert use_ocr is True
        return ExtractedDocument(
            path,
            [(1, "<table><tr><td>项目</td><td>12</td></tr></table>")],
            ocr_used=True,
        )

    process_document_job(job_dir, None, extract_fn=fake_extract)
    assert (job_dir / "01_Extracted_Text" / "extracted_text.md").is_file()
    saved_text = (job_dir / "01_Extracted_Text" / "extracted_text.md").read_text(encoding="utf-8")
    assert "<td>12</td>" in saved_text
    assert "项目" in saved_text


def test_older_document_job_keeps_its_flat_output_paths(tmp_path: Path):
    from runtime.documents import create_document_job, load_document_job, process_document_job, _save_document_state

    source = tmp_path / "old-notes.txt"
    source.write_text("A small source document.", encoding="utf-8")
    job_dir = create_document_job(source, ["summary"], False, tmp_path / "document_outputs")
    state = load_document_job(job_dir)
    state.pop("output_layout_version")
    _save_document_state(job_dir, state)

    process_document_job(job_dir, lambda _prompt: json.dumps({"summary": "An old-layout result."}))
    assert (job_dir / "summary.md").read_text(encoding="utf-8") == "An old-layout result."
    assert not (job_dir / "START_HERE.txt").exists()
    assert not (job_dir / "02_AI_Results" / "summary.md").exists()


@pytest.mark.parametrize(
    ("suffix", "use_ocr", "expected_relative"),
    [
        (".txt", False, "01_Extracted_Text/extracted_text.txt"),
        (".png", True, "01_Extracted_Text/extracted_text.md"),
    ],
)
def test_categorized_v2_job_keeps_its_original_text_extension(
    tmp_path: Path, suffix: str, use_ocr: bool, expected_relative: str,
):
    from runtime.documents import (
        ExtractedDocument, _save_document_state, create_document_job, load_document_job,
        process_document_job,
    )

    source = tmp_path / f"older{suffix}"
    source.write_text("Older source.", encoding="utf-8") if suffix == ".txt" else source.write_bytes(b"mock image")
    job_dir = create_document_job(source, ["verbatim"], use_ocr, tmp_path / "document_outputs")
    state = load_document_job(job_dir)
    state["output_layout_version"] = 2
    _save_document_state(job_dir, state)

    def fake_extract(path: Path, *, use_ocr: bool = False) -> ExtractedDocument:
        return ExtractedDocument(path, [(1, "Recovered older job.")], ocr_used=use_ocr)

    process_document_job(job_dir, None, extract_fn=fake_extract)
    assert (job_dir / expected_relative).is_file()
    assert not (job_dir / "START_HERE.txt").exists()


def test_pdf_table_extraction_writes_markdown_and_keeps_surrounding_text(tmp_path: Path):
    from runtime.documents import create_document_job, process_document_job

    source = Path(__file__).parent / "fixtures" / "structured_table.pdf"
    job_dir = create_document_job(source, ["verbatim"], False, tmp_path / "document_outputs")
    process_document_job(job_dir, None)

    output = job_dir / "01_Extracted_Text" / "extracted_text.md"
    assert output.is_file()
    text = output.read_text(encoding="utf-8")
    assert "| Product | Revenue | Year |" in text
    assert "| Alpha | 1,200 | 2025 |" in text
    assert "End of report." in text
    assert not (job_dir / "START_HERE.txt").exists()




def test_docx_keeps_heading_emphasis_paragraph_order_and_table(tmp_path: Path):
    from docx import Document

    source = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Quarterly report", level=1)
    paragraph = document.add_paragraph("Revenue increased; ")
    paragraph.add_run("important figure").bold = True
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Product"
    table.rows[0].cells[1].text = "Revenue"
    row = table.add_row().cells
    row[0].text, row[1].text = "Alpha", "1200"
    document.add_paragraph("End of report.")
    document.save(source)

    extracted = extract_document(source)
    text = extracted.text

    assert "# Quarterly report" in text
    assert "**important figure**" in text
    assert "| Product | Revenue |" in text
    assert "| Alpha | 1200 |" in text
    assert text.index("| Alpha | 1200 |") < text.index("End of report.")
    assert "[Page 1]" not in text


def test_csv_keeps_zero_and_quoted_multiline_cells(tmp_path: Path):
    source = tmp_path / "values.csv"
    source.write_text('Name,Amount,Notes\r\nItem,0,"line one\nline two"\r\n', encoding="utf-8")

    text = extract_document(source).text

    assert "| Item | 0 | line one<br>line two |" in text


def test_xlsx_keeps_sheet_names_zero_values_and_cell_rows(tmp_path: Path):
    from openpyxl import Workbook

    source = tmp_path / "figures.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Annual figures"
    worksheet.append(["Item", "Amount"])
    worksheet.append(["Unassigned", 0])
    workbook.save(source)
    workbook.close()

    extracted = extract_document(source)
    text = extracted.text

    assert "## Worksheet: Annual figures" in text
    assert "| Item | Amount |" in text
    assert "| Unassigned | 0 |" in text
    assert any("Excel cell values" in warning for warning in extracted.warnings)


def test_legacy_document_checkpoint_is_still_loadable(tmp_path: Path):
    import json
    from runtime.documents import load_document_job, list_resumable_document_jobs

    output_root = tmp_path / "outputs"
    job_dir = output_root / "old_job"
    job_dir.mkdir(parents=True)
    legacy_state = {
        "version": 1, "source": str(tmp_path / "old.txt"), "selected": ["verbatim"],
        "status": "pending", "updated_at": "2026-01-01T00:00:00+00:00",
    }
    (job_dir / ".taskrunner-study.json").write_text(json.dumps(legacy_state), encoding="utf-8")

    assert load_document_job(job_dir)["status"] == "pending"
    resumable = list_resumable_document_jobs(output_root)
    assert len(resumable) == 1
    assert resumable[0]["job_dir"] == str(job_dir.resolve())


def test_legacy_start_here_file_is_removed_when_job_completes(tmp_path: Path):
    from runtime.documents import create_document_job, process_document_job

    source = tmp_path / "notes.txt"
    source.write_text("A short note.", encoding="utf-8")
    job_dir = create_document_job(source, ["verbatim"], False, tmp_path / "document_outputs")
    legacy_index = job_dir / "START_HERE.txt"
    legacy_index.write_text("Old instructions", encoding="utf-8")

    process_document_job(job_dir, None)

    assert not legacy_index.exists()
    assert (job_dir / "01_Extracted_Text" / "extracted_text.md").is_file()
